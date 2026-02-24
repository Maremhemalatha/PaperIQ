import streamlit as st
import pdfplumber
import docx as docx_lib
import re
import json
import numpy as np
import heapq
import os
from datetime import datetime
from typing import Dict, List, Tuple, Optional
from textblob import TextBlob
import plotly.graph_objects as go
from fpdf import FPDF

st.set_page_config(
    page_title="PaperIQ - AI Research Analyzer",
    page_icon=" ",
    layout="wide",
    initial_sidebar_state="expanded"
)
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Crimson+Pro:wght@400;600;700&family=IBM+Plex+Mono:wght@400;500&family=Playfair+Display:wght@700;900&display=swap');

    .main { background-color: #faf8f3; }
    h1 { font-family: 'Playfair Display', serif; color: #1a1a2e; font-weight: 900; font-size: 3rem; letter-spacing: -1px; }
    h2, h3 { font-family: 'Playfair Display', serif; color: #c7522a; }

    .stButton>button {
        background-color: #c7522a; color: white;
        font-family: 'IBM Plex Mono', monospace; font-weight: 500;
        border: none; padding: 0.75rem 2rem;
        text-transform: uppercase; letter-spacing: 1px;
        transition: all 0.3s ease;
    }
    .stButton>button:hover { background-color: #1a1a2e; transform: translateX(4px); }

    .keyword-badge {
        display: inline-block; background: #fbf4e8;
        border: 1px solid #e0ddd4; padding: 0.4rem 1rem; margin: 0.25rem;
        font-family: 'IBM Plex Mono', monospace; font-size: 0.85rem; border-radius: 4px;
    }
    .domain-badge {
        display: inline-block; background: #2a5298; color: white;
        padding: 0.8rem 1.5rem; font-family: 'IBM Plex Mono', monospace;
        text-transform: uppercase; letter-spacing: 1px;
        border-radius: 4px; margin-top: 1rem;
    }
    .metric-card {
        background: white;
        color: black;
        padding: 1rem;
        border-radius: 4px;
        text-align: center;
    }
    .summary-card {
        background: #f0f4ff; border: 1px solid #c0cfe8;
        border-left: 6px solid #2a5298; padding: 1.2rem 1.5rem;
        border-radius: 4px; margin-bottom: 0.75rem;
        font-family: 'Crimson Pro', serif; font-size: 1.05rem; line-height: 1.7;
    }
    .paper-summary-card {
        background: #fff8f0; border: 1px solid #e8cfa0;
        border-left: 6px solid #c7522a; padding: 1.4rem 1.6rem;
        border-radius: 4px; margin-bottom: 1rem;
        font-family: 'Crimson Pro', serif; font-size: 1.1rem; line-height: 1.8;
    }
    .quality-label { font-family: 'IBM Plex Mono', monospace; font-size: 0.78rem; color: #555; margin-bottom: 2px; }
    .quality-bar-wrap { background: #e0ddd4; border-radius: 99px; height: 10px; width: 100%; margin: 4px 0 12px 0; }
    .quality-bar-fill { height: 10px; border-radius: 99px; background: linear-gradient(90deg, #2a5298, #c7522a); }
    .score-chip { display: inline-block; padding: 0.2rem 0.65rem; border-radius: 99px; font-family: 'IBM Plex Mono', monospace; font-size: 0.78rem; font-weight: 600; margin-left: 6px; }
    .score-high   { background: #d4edda; color: #155724; }
    .score-medium { background: #fff3cd; color: #856404; }
    .score-low    { background: #f8d7da; color: #721c24; }
</style>
""", unsafe_allow_html=True)

SECTION_IMPORTANCE = {
    "abstract":        {"importance": "high",   "target_words": 120, "style": "concise overview"},
    "introduction":    {"importance": "high",   "target_words": 150, "style": "motivation and context"},
    "methodology":     {"importance": "high",   "target_words": 160, "style": "technical approach with key steps"},
    "methods":         {"importance": "high",   "target_words": 160, "style": "technical approach with key steps"},
    "results":         {"importance": "high",   "target_words": 150, "style": "key findings and metrics"},
    "conclusion":      {"importance": "high",   "target_words": 120, "style": "main takeaways and impact"},
    "related work":    {"importance": "medium", "target_words": 100, "style": "key prior work and gaps"},
    "discussion":      {"importance": "medium", "target_words": 110, "style": "interpretation of results"},
    "evaluation":      {"importance": "medium", "target_words": 100, "style": "evaluation setup and outcome"},
    "implementation":  {"importance": "medium", "target_words":  90, "style": "implementation details"},
    "future work":     {"importance": "medium", "target_words":  80, "style": "future directions"},
    "acknowledgments": {"importance": "low",    "target_words":  40, "style": "brief mention"},
    "references":      {"importance": "low",    "target_words":  40, "style": "brief mention"},
    "keywords":        {"importance": "low",    "target_words":  30, "style": "list"},
}

def get_section_config(section_name: str) -> dict:
    name_lower = section_name.lower()
    for key, cfg in SECTION_IMPORTANCE.items():
        if key in name_lower:
            return cfg
    return {"importance": "medium", "target_words": 100, "style": "key points"}

@st.cache_resource
def load_hf_summarizer():
    try:
        from transformers import pipeline
        return pipeline("summarization", model="facebook/bart-large-cnn")
    except Exception:
        return None

class HuggingFaceSummarizer:
    def __init__(self):
        self.summarizer = load_hf_summarizer()

    def is_available(self) -> bool:
        return self.summarizer is not None

    def summarize_section(self, section_name: str, section_text: str) -> Optional[str]:
        if not self.summarizer:
            return None
        cfg        = get_section_config(section_name)
        target     = cfg["target_words"]
        max_tokens = min(target + 50, 200)
        min_tokens = max(30, target // 2)
        try:
            text_input = section_text[:1024]
            if len(text_input.split()) < 30:
                return None
            result = self.summarizer(
                text_input,
                max_length=max_tokens,
                min_length=min_tokens,
                do_sample=False,
            )
            return result[0]["summary_text"].strip()
        except Exception:
            return None

    def summarize_full_paper(self, sections: Dict[str, str]) -> Optional[str]:
        if not self.summarizer:
            return None
        key_order = ["Abstract", "Introduction", "Methodology", "Results", "Conclusion"]
        parts = []
        for key in key_order:
            for sname, scontent in sections.items():
                if key.lower() in sname.lower():
                    parts.append(scontent[:400])
                    break
        if not parts:
            parts = [v[:400] for v in list(sections.values())[:4]]
        combined = " ".join(parts)[:1024]
        try:
            if len(combined.split()) < 30:
                return None
            result = self.summarizer(combined, max_length=220, min_length=80, do_sample=False)
            return result[0]["summary_text"].strip()
        except Exception:
            return None
          
class SummaryQualityScorer:
    FILLER_PHRASES = [
        "this section", "the authors", "it is noted", "in this paper",
        "as mentioned", "it can be seen", "it is important to note",
        "this study discusses", "this paper discusses",
    ]

    def coherence_score(self, summary: str) -> float:
        if not summary or len(summary) < 20:
            return 0.0
        sentences  = re.split(r'(?<=[.!?])\s+', summary.strip())
        score      = 70.0
        word_count = len(summary.split())
        if word_count < 30:
            score -= 20
        elif word_count < 50:
            score -= 10
        starters = [s.split()[0].lower() for s in sentences if s.split()]
        score -= (len(starters) - len(set(starters))) * 5
        summary_lower = summary.lower()
        score -= sum(1 for fp in self.FILLER_PHRASES if fp in summary_lower) * 8
        n = len(sentences)
        if n >= 3: score += 10
        if n >= 5: score += 5
        return max(0.0, min(100.0, score))

    def readability_score(self, summary: str) -> float:
        if not summary or len(summary) < 10:
            return 0.0
        words     = re.findall(r'\b\w+\b', summary)
        sentences = re.split(r'(?<=[.!?])\s+', summary.strip())
        n_words   = len(words)
        n_sents   = max(len(sentences), 1)

        def syllables(word: str) -> int:
            word  = word.lower()
            count = len(re.findall(r'[aeiou]+', word))
            if word.endswith('e') and len(word) > 2:
                count = max(count - 1, 1)
            return max(count, 1)

        n_syllables = sum(syllables(w) for w in words)
        if n_words == 0:
            return 0.0
        raw = 206.835 - 1.015 * (n_words / n_sents) - 84.6 * (n_syllables / n_words)
        if 30 <= raw <= 60:   quality = 85 + (raw - 30) / 30 * 15
        elif raw > 60:        quality = max(40, 100 - (raw - 60) * 1.5)
        else:                 quality = max(20, 60 + raw)
        return round(min(100.0, max(0.0, quality)), 1)

    def length_optimality_score(self, summary: str, section_name: str) -> float:
        target = get_section_config(section_name)["target_words"]
        actual = len(summary.split())
        if actual == 0:
            return 0.0
        ratio = actual / target
        if 0.8 <= ratio <= 1.2: return 100.0
        if 0.6 <= ratio <= 1.5: return 75.0
        if 0.4 <= ratio <= 2.0: return 50.0
        return 25.0

    def overall_score(self, summary: str, section_name: str) -> dict:
        c       = self.coherence_score(summary)
        r       = self.readability_score(summary)
        l       = self.length_optimality_score(summary, section_name)
        overall = round(c * 0.4 + r * 0.35 + l * 0.25, 1)

        def label(v):
            if v >= 75: return "high"
            if v >= 50: return "medium"
            return "low"

        return {
            "coherence": round(c, 1), "readability": round(r, 1),
            "length_opt": round(l, 1), "overall": overall,
            "coherence_label": label(c), "readability_label": label(r),
            "length_label": label(l), "overall_label": label(overall),
        }

def render_quality_bar(label: str, score: float, level: str) -> str:
    return f"""
    <div class="quality-label">{label}
        <span class="score-chip score-{level}">{score}/100</span>
    </div>
    <div class="quality-bar-wrap">
        <div class="quality-bar-fill" style="width:{score}%"></div>
    </div>
    """
class PaperAnalyzer:
    def extract_text_from_pdf(self, pdf_file) -> Tuple[str, int]:
        try:
            with pdfplumber.open(pdf_file) as pdf:
                text = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
                return text, len(pdf.pages)
        except Exception as e:
            st.error(f"Error extracting PDF: {e}")
            return "", 0

    def extract_text_from_docx(self, docx_file) -> Tuple[str, int]:
        try:
            doc        = docx_lib.Document(docx_file)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            text       = "\n".join(paragraphs)
            pages      = max(1, len(text.split()) // 250)
            return text, pages
        except Exception as e:
            st.error(f"Error extracting DOCX: {e}")
            return "", 0

    def extract_text_from_txt(self, txt_file) -> Tuple[str, int]:
        try:
            text  = txt_file.getvalue().decode("utf-8")
            pages = max(1, len(text.split()) // 250)
            return text, pages
        except Exception as e:
            st.error(f"Error reading TXT: {e}")
            return "", 0

    def clean_text(self, text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'[^\w\s.,;:!()\-"]', '', text)
        return text.strip()

    def count_words(self, text: str) -> int:
        return len(re.findall(r'\b\w+\b', text))

    def identify_sections(self, text: str) -> Dict[str, str]:
        sections:   Dict[str, str]  = {}
        candidates: Dict[str, list] = {}

        for num, title, content in re.findall(
            r'(?:^|\n)\s*(\d+)\.\s+([A-Z][A-Za-z\s&-]+?)(?:\n|:)(.*?)(?=\n\s*\d+\.\s+[A-Z]|\Z)',
            text, re.DOTALL
        ):
            key  = f"{num}. {title.strip()}"
            body = content.strip()
            if body and len(body) > 80:
                candidates.setdefault(num, []).append((key, body))

        for num, matches in candidates.items():
            best_key, best_body = max(matches, key=lambda x: len(x[1]))
            sections[best_key]  = best_body[:3000]

        unnumbered = {
            'Abstract':        r'(?i)(?:^|\n)\s*(?:abstract|summary)\s*[:\n](.*?)(?=\n\s*(?:introduction|keywords|1\.|background|\Z))',
            'Keywords':        r'(?i)(?:^|\n)\s*keywords?\s*[:\-]?\s*([^\n]+)',
            'Introduction':    r'(?i)(?:^|\n)\s*introduction\s*[:\n](.*?)(?=\n\s*(?:2\.|methodology|methods|related work|literature|background|\Z))',
            'Related Work':    r'(?i)(?:^|\n)\s*(?:related\s+work|literature\s+review|background)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|methodology|methods|\Z))',
            'Methodology':     r'(?i)(?:^|\n)\s*(?:methodology|methods|materials\s+and\s+methods|approach|experimental\s+setup)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|results|experiments|implementation|\Z))',
            'Results':         r'(?i)(?:^|\n)\s*(?:results|findings|experiments|experimental\s+results)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|discussion|conclusion|analysis|\Z))',
            'Discussion':      r'(?i)(?:^|\n)\s*discussion\s*[:\n](.*?)(?=\n\s*(?:\d+\.|conclusion|references|\Z))',
            'Conclusion':      r'(?i)(?:^|\n)\s*(?:conclusion|conclusions|concluding\s+remarks)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|references|acknowledgment|\Z))',
            'Future Work':     r'(?i)(?:^|\n)\s*(?:future\s+work|future\s+directions|future\s+research)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|references|acknowledgment|\Z))',
            'Acknowledgments': r'(?i)(?:^|\n)\s*acknowledgments?\s*[:\n](.*?)(?=\n\s*(?:\d+\.|references|\Z))',
            'References':      r'(?i)(?:^|\n)\s*references\s*[:\n](.*?)(?=\Z)',
        }

        for name, pattern in unnumbered.items():
            if name not in sections:
                m = re.search(pattern, text, re.DOTALL)
                if m:
                    body = m.group(1).strip()
                    if body and len(body) > 10:
                        sections[name] = body[:3000]

        for num, title, content in re.findall(
            r'(?:^|\n)\s*(\d+\.\d+)\s+([A-Z][A-Za-z\s&-]+?)(?:\n|:)(.*?)(?=\n\s*\d+\.\d+\s+[A-Z]|\n\s*\d+\.\s+[A-Z]|\Z)',
            text, re.DOTALL
        ):
            key  = f"{num} {title.strip()}"
            body = content.strip()
            if body and len(body) > 20:
                sections[key] = body[:2000]

        alternatives = {
            'Data Collection':    r'(?i)(?:^|\n)\s*(?:data\s+collection|data\s+gathering)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Data Analysis':      r'(?i)(?:^|\n)\s*(?:data\s+analysis|statistical\s+analysis)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Experimental Setup': r'(?i)(?:^|\n)\s*(?:experimental\s+setup|experimental\s+design)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Implementation':     r'(?i)(?:^|\n)\s*implementation\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Evaluation':         r'(?i)(?:^|\n)\s*evaluation\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Limitations':        r'(?i)(?:^|\n)\s*limitations\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
            'Contributions':      r'(?i)(?:^|\n)\s*(?:contributions|our\s+contributions)\s*[:\n](.*?)(?=\n\s*(?:\d+\.|\w+\s+\w+|\Z))',
        }

        for name, pattern in alternatives.items():
            if name not in sections:
                m = re.search(pattern, text, re.DOTALL)
                if m:
                    body = m.group(1).strip()
                    if body and len(body) > 20:
                        sections[name] = body[:2000]

        return sections or {"Note": "No clear sections found. The paper may use non-standard formatting."}

    def extract_keywords(self, text: str) -> List[str]:
        keywords: set = set()
        m = re.search(r'(?i)keywords?\s*[:\-]?\s*([^\n]+)', text)
        if m:
            keywords.update(kw.strip() for kw in re.split(r'[,;·•]', m.group(1)) if kw.strip())
        if len(keywords) < 5:
            terms = re.findall(r'\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+)*\b', text)
            freq: Dict[str, int] = {}
            for t in terms:
                if len(t) > 4:
                    freq[t] = freq.get(t, 0) + 1
            keywords.update(t for t, _ in sorted(freq.items(), key=lambda x: -x[1])[:10])
        return list(keywords)[:12]

    def classify_domain(self, text: str, keywords: List[str]) -> str:
        domains = {
            'Computer Science & AI':     ['algorithm', 'machine learning', 'artificial intelligence', 'neural network', 'deep learning', 'computer', 'software', 'AI', 'ML', 'computational'],
            'Medical & Health Sciences': ['patient', 'clinical', 'medical', 'health', 'disease', 'treatment', 'therapy', 'diagnosis', 'hospital', 'medicine'],
            'Engineering':               ['engineering', 'design', 'system', 'control', 'optimization', 'mechanical', 'electrical', 'civil', 'structural'],
            'Physics':                   ['quantum', 'energy', 'physics', 'particle', 'electromagnetic', 'thermodynamic', 'mechanics', 'radiation'],
            'Chemistry':                 ['chemical', 'molecule', 'synthesis', 'reaction', 'compound', 'polymer', 'catalyst', 'organic', 'inorganic'],
            'Biology':                   ['biological', 'cell', 'gene', 'protein', 'DNA', 'organism', 'evolution', 'species', 'ecology'],
            'Mathematics':               ['mathematical', 'theorem', 'proof', 'equation', 'formula', 'matrix', 'optimization', 'statistics'],
            'Social Sciences':           ['social', 'society', 'behavior', 'psychology', 'economic', 'cultural', 'political', 'demographic'],
            'Environmental Science':     ['environment', 'climate', 'pollution', 'sustainability', 'ecosystem', 'conservation', 'renewable'],
        }
        combined = (text + ' ' + ' '.join(keywords)).lower()
        scores   = {d: sum(1 for kw in kws if kw.lower() in combined) for d, kws in domains.items()}
        best     = max(scores.items(), key=lambda x: x[1])
        return best[0] if best[1] > 0 else "Interdisciplinary Research"

    def extractive_summary(self, section_text: str, section_name: str = "") -> str:
        if not section_text:
            return ""
        target    = get_section_config(section_name)["target_words"]
        sentences = re.split(r'(?<=[.!?])\s+', section_text.strip())
        if not sentences:
            return section_text[:500]
        all_words = re.findall(r'\b\w{4,}\b', section_text.lower())
        freq: Dict[str, int] = {}
        for w in all_words:
            freq[w] = freq.get(w, 0) + 1

        def score(s: str) -> float:
            words = re.findall(r'\b\w{4,}\b', s.lower())
            return sum(freq.get(w, 0) for w in words) / max(len(words), 1)

        ranked      = sorted(range(len(sentences)), key=lambda i: -score(sentences[i]))
        picked: set = set()
        word_count  = 0
        for idx in ranked:
            if word_count >= target:
                break
            picked.add(idx)
            word_count += len(sentences[idx].split())
        return " ".join(s for i, s in enumerate(sentences) if i in picked).strip()

def calculate_readability(text: str) -> float:
    sentences = text.count('.') + text.count('!') + text.count('?')
    words     = len(text.split())
    syllables = int(words * 1.5)
    if sentences == 0 or words == 0:
        return 0.0
    score = 206.835 - 1.015 * (words / sentences) - 84.6 * (syllables / words)
    return max(0.0, min(100.0, score))

def analyze_writing_quality(text: str) -> dict:
    blob           = TextBlob(text)
    sentences      = blob.sentences
    words          = blob.words
    word_count     = len(words)
    sentence_count = len(sentences)

    if sentence_count == 0 or word_count == 0:
        return None

    avg_sentence_len = float(np.mean([len(s.words) for s in sentences]))
    avg_word_len     = float(np.mean([len(w) for w in words]))
    sentiment        = blob.sentiment.polarity

    language_score   = min(100.0, (avg_sentence_len * 1.5) + (avg_word_len * 5) + (50 + sentiment * 20))

    transitions      = ["however", "therefore", "thus", "consequently", "furthermore", "meanwhile"]
    transition_count = sum(text.lower().count(t) for t in transitions)
    coherence_score  = min(100.0, (transition_count * 4) + (sentence_count * 0.1) + 40)

    reasoning_kws   = ["because", "since", "implies", "due to", "as a result", "evidence"]
    reasoning_count = sum(text.lower().count(k) for k in reasoning_kws)
    reasoning_score = min(100.0, (reasoning_count * 6) + 30)

    complex_words = [w for w in words if len(w) > 6]
    lexical_score = min(100.0, (len(complex_words) / word_count) * 300) if word_count else 0.0
    readability   = calculate_readability(text)

    composite = (language_score * 0.3 + coherence_score * 0.2 +
                 reasoning_score * 0.2 + lexical_score * 0.15 + readability * 0.15)

    long_sentences  = [s.raw for s in sentences if len(s.words) > 30]
    vocab_diversity = round(len(set(w.lower() for w in words)) / word_count, 2) if word_count else 0
    return {
        "scores": {
            "Language":       round(language_score, 1),
            "Coherence":      round(coherence_score, 1),
            "Reasoning":      round(reasoning_score, 1),
            "Sophistication": round(lexical_score, 1),
            "Readability":    round(readability, 1),
            "Composite":      round(composite, 1),
        },
        "stats": {
            "word_count":         word_count,
            "sentence_count":     sentence_count,
            "avg_sentence_len":   round(avg_sentence_len, 2),
            "avg_word_len":       round(avg_word_len, 2),
            "vocab_diversity":    vocab_diversity,
            "complex_word_ratio": round(len(complex_words) / word_count, 2) if word_count else 0,
        },
        "sentiment":      round(sentiment, 2),
        "long_sentences": long_sentences,
        "blob":           blob,
    }

def create_pdf_report(filename: str, results: dict, writing: dict, section_summaries: dict, paper_summary: str) -> bytes:
    pdf  = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    safe = lambda s: str(s).encode('latin-1', 'replace').decode('latin-1')

    pdf.cell(0, 10, safe("PaperIQ Analysis Report"), ln=1, align='C')
    pdf.set_font("Arial", size=11)
    pdf.cell(0, 8, safe(f"File: {filename}"), ln=1)
    pdf.cell(0, 8, safe(f"Domain: {results.get('domain', '')}"), ln=1)
    pdf.cell(0, 8, safe(f"Pages: {results.get('num_pages','')}  |  Words: {results.get('word_count','')}  |  Sections: {len(results.get('sections',{}))}"), ln=1)
    pdf.ln(4)

    if writing:
        scores = writing["scores"]
        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, "Writing Quality Scores", ln=1)
        pdf.set_font("Arial", size=11)
        pdf.cell(0, 8, safe(f"Composite Score: {scores['Composite']}/100"), ln=1)
        for k, v in scores.items():
            if k != "Composite":
                pdf.cell(0, 7, safe(f"  {k}: {v}/100"), ln=1)
        pdf.ln(3)

        stats = writing["stats"]
        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, "Text Statistics", ln=1)
        pdf.set_font("Arial", size=11)
        for k, v in stats.items():
            pdf.cell(0, 7, safe(f"  {k.replace('_',' ').title()}: {v}"), ln=1)
        pdf.ln(3)

        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, safe(f"Sentiment: {writing['sentiment']} ({'Positive' if writing['sentiment'] > 0 else 'Neutral/Negative'})"), ln=1)
        pdf.ln(3)

    if paper_summary:
        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, "Paper Summary", ln=1)
        pdf.set_font("Arial", size=11)
        pdf.multi_cell(0, 7, safe(paper_summary))
        pdf.ln(3)

    if section_summaries:
        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, "Section Summaries", ln=1)
        pdf.set_font("Arial", size=10)
        for title, sdata in section_summaries.items():
            text = sdata["text"] if isinstance(sdata, dict) else sdata
            pdf.set_font("Arial", 'B', 10)
            pdf.cell(0, 7, safe(f"[{title}]"), ln=1)
            pdf.set_font("Arial", size=10)
            pdf.multi_cell(0, 6, safe(text))
            pdf.ln(2)

    if results.get("keywords"):
        pdf.set_font("Arial", 'B', 13)
        pdf.cell(0, 9, "Keywords", ln=1)
        pdf.set_font("Arial", size=11)
        pdf.multi_cell(0, 7, safe(", ".join(results["keywords"])))

    return pdf.output(dest='S').encode('latin-1')

def initialize_session_state():
    defaults = {
        'authenticated': False, 'user_role': None, 'user_email': None,
        'analysis_complete': False, 'analysis_results': None,
        'section_summaries': {}, 'paper_summary': None, 'writing_quality': None,
    }
    for key, val in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = val

def validate_email(email: str) -> bool:
    return bool(re.match(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$', email))

def login_page():
    st.title(" PaperIQ")
    st.markdown("### AI-Powered Research Insight Analyzer")
    st.markdown("---")
    _, col, _ = st.columns([1, 2, 1])
    with col:
        st.markdown("#### Welcome to PaperIQ")
        st.write("Enter your details to begin analyzing research papers")
        st.markdown("<br>", unsafe_allow_html=True)
        email    = st.text_input(" Email Address", placeholder="your.email@example.com")
        password = st.text_input(" Password", type="password", placeholder="Enter your password (min 6 chars)")
        st.markdown("<br>", unsafe_allow_html=True)
        st.write("**Select your role:**")
        c1, c2        = st.columns(2)
        selected_role = None
        with c1:
            if st.button(" Student",    use_container_width=True): selected_role = "Student"
        with c2:
            if st.button(" Researcher", use_container_width=True): selected_role = "Researcher"
        if selected_role:
            if not email:                   st.error("Please enter your email address")
            elif not validate_email(email): st.error("Please enter a valid email address")
            elif not password:              st.error("Please enter your password")
            elif len(password) < 6:         st.error("Password must be at least 6 characters")
            else:
                st.session_state.update(user_role=selected_role, user_email=email, authenticated=True)
                st.success(f"Logged in as {selected_role}")
                st.rerun()

def upload_page():
    with st.sidebar:
        st.markdown(f"###  {st.session_state.user_role}")
        st.markdown(f" {st.session_state.user_email}")
        st.markdown("---")
        st.markdown("#### PaperIQ Features")
        st.write("""
        - PDF, DOCX, TXT support
        - AI summaries (HuggingFace BART / Extractive)
        - Writing quality scoring
        - Radar & bar chart visualizations
        - Sentiment analysis
        - Long sentence detection
        - Vocabulary suggestions
        - Keywords & domain classification
        - PDF report download
        - JSON export
        """)
        st.markdown("---")
        if st.button(" Logout"):
            for k in ['authenticated', 'user_role', 'user_email', 'analysis_complete',
                      'analysis_results', 'section_summaries', 'paper_summary', 'writing_quality']:
                st.session_state[k] = (False if k in ('authenticated', 'analysis_complete')
                                       else ({} if k == 'section_summaries' else None))
            st.rerun()

    st.title(" PaperIQ Research Analyzer")
    st.markdown(f"**Logged in as:** {st.session_state.user_role} ({st.session_state.user_email})")
    st.markdown("---")

    if not st.session_state.analysis_complete:
        st.markdown("###  Upload Research Paper")
        st.write("Upload a PDF, DOCX, or TXT file for comprehensive AI-powered analysis")
        uploaded_file = st.file_uploader("Choose a file", type=["pdf", "docx", "txt"])
        if uploaded_file:
            st.success(f" File uploaded: {uploaded_file.name}")
            if st.button(" Analyze Paper", use_container_width=True):
                analyze_paper(uploaded_file)
    else:
        display_analysis()

def analyze_paper(pdf_file):
    analyzer = PaperAnalyzer()
    scorer   = SummaryQualityScorer()

    with st.spinner(" Extracting text…"):
        fname = pdf_file.name.lower()
        if fname.endswith(".pdf"):
            text, num_pages = analyzer.extract_text_from_pdf(pdf_file)
        elif fname.endswith(".docx"):
            text, num_pages = analyzer.extract_text_from_docx(pdf_file)
        else:
            text, num_pages = analyzer.extract_text_from_txt(pdf_file)
        if not text:
            st.error("Could not extract text. Please try another file.")
            return
        st.info(f" Extracted {num_pages} pages")

    with st.spinner(" Processing text…"):
        cleaned    = analyzer.clean_text(text)
        word_count = analyzer.count_words(cleaned)

    with st.spinner(" Identifying sections…"):
        sections = analyzer.identify_sections(text)
        st.info(f" Found {len(sections)} sections")

    with st.spinner(" Extracting keywords…"):
        keywords = analyzer.extract_keywords(text)

    with st.spinner(" Classifying domain…"):
        domain = analyzer.classify_domain(text, keywords)

    with st.spinner(" Analyzing writing quality…"):
        writing_quality = analyze_writing_quality(cleaned)

    section_summaries: Dict[str, dict] = {}
    paper_summary: Optional[str]       = None
    SKIP = {"references", "keywords", "acknowledgments"}
    summarizable = {
        k: v for k, v in sections.items()
        if not any(s in k.lower() for s in SKIP) and len(v.split()) > 40
    }

    hf = HuggingFaceSummarizer()
    if hf.is_available():
        total    = max(len(summarizable), 1)
        progress = st.progress(0, text=" Generating summaries (HuggingFace BART)…")
        for i, (sname, scontent) in enumerate(summarizable.items()):
            progress.progress((i + 1) / total, text=f" Summarising: {sname}")
            hf_text = hf.summarize_section(sname, scontent)
            if hf_text:
                section_summaries[sname] = {"text": hf_text, "quality": scorer.overall_score(hf_text, sname), "source": "huggingface"}
            else:
                fallback = analyzer.extractive_summary(scontent, sname)
                section_summaries[sname] = {"text": fallback, "quality": scorer.overall_score(fallback, sname), "source": "extractive"}
        progress.empty()
        with st.spinner(" Generating paper-level summary…"):
            paper_summary = hf.summarize_full_paper(sections)
    else:
        st.warning(" HuggingFace not available — using extractive summaries. Run: pip install transformers torch")
        for sname, scontent in summarizable.items():
            fallback = analyzer.extractive_summary(scontent, sname)
            section_summaries[sname] = {"text": fallback, "quality": scorer.overall_score(fallback, sname), "source": "extractive"}

    st.session_state.section_summaries = section_summaries
    st.session_state.paper_summary     = paper_summary
    st.session_state.writing_quality   = writing_quality
    st.session_state.analysis_results  = {
        'filename': pdf_file.name, 'num_pages': num_pages, 'word_count': word_count,
        'sections': sections, 'keywords': keywords, 'domain': domain,
        'analyzed_date': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        'user_email': st.session_state.user_email, 'user_role': st.session_state.user_role,
    }
    st.session_state.analysis_complete = True
    st.success(f" Analysis complete! Found {len(sections)} sections")
    st.rerun()

def display_analysis():
    results           = st.session_state.analysis_results
    section_summaries = st.session_state.get("section_summaries", {})
    paper_summary     = st.session_state.get("paper_summary")
    writing           = st.session_state.get("writing_quality")

    st.markdown("###  Analysis Results")
    st.markdown(f"**Paper:** {results['filename']}")
    st.markdown("---")

    c1, c2, c3, c4 = st.columns(4)
    for col, icon, label, val in [
        (c1, " ", "Pages",    results['num_pages']),
        (c2, " ", "Words",    f"{results['word_count']:,}"),
        (c3, " ", "Sections", len(results['sections'])),
        (c4, " ", "Keywords", len(results['keywords'])),
    ]:
        with col:
            st.markdown(f'<div class="metric-card"><h4>{icon} {label}</h4><h2>{val}</h2></div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    if writing:
        scores = writing["scores"]
        wc1, wc2, wc3, wc4, wc5, wc6 = st.columns(6)
        for col, key in zip([wc1, wc2, wc3, wc4, wc5, wc6], ["Composite", "Language", "Coherence", "Reasoning", "Sophistication", "Readability"]):
            with col:
                st.metric(key, f"{scores[key]}/100")

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("###  Domain Classification")
    st.markdown(f'<div class="domain-badge">{results["domain"]}</div>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("###  Extracted Keywords")
    st.markdown('<div>' + ''.join(f'<span class="keyword-badge">{kw}</span>' for kw in results['keywords']) + '</div>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)

    if paper_summary:
        st.markdown("###  Paper Summary")
        st.markdown(f'<div class="paper-summary-card">{paper_summary}</div>', unsafe_allow_html=True)
        st.markdown("<br>", unsafe_allow_html=True)

    has_hf   = any(v.get("source") == "huggingface" for v in section_summaries.values())
    ai_badge = "HuggingFace BART" if has_hf else " Extractive"
    st.markdown(f"### All Sections ({len(results['sections'])} Found) — {ai_badge}")

    tab_sections, tab_viz, tab_issues, tab_suggestions, tab_stats, tab_sentiment = st.tabs([
        " Section Summaries", " Visualizations", " Issues", " Suggestions", " Detailed Metrics", " Sentiment"
    ])

    with tab_sections:
        st.info("Each section shows full text and a summary with quality scores. Click to expand.")
        numbered, subsect, other = {}, {}, {}
        for name, content in results['sections'].items():
            if re.match(r'^\d+\.\s', name):      numbered[name] = content
            elif re.match(r'^\d+\.\d+\s', name): subsect[name]  = content
            else:                                  other[name]    = content

        def render_section(sname: str, content: str):
            summary_data = section_summaries.get(sname)
            with st.expander(sname, expanded=False):
                if not content or content == "Section not found or could not be extracted.":
                    st.warning("Content not available")
                    return
                tab_raw, tab_sum = st.tabs([" Full Text", " Summary & Quality"])
                with tab_raw:
                    st.write(content)
                    st.caption(f"Word count: {len(content.split())} words")
                with tab_sum:
                    if not summary_data:
                        st.info("No summary available for this section.")
                        return
                    source = summary_data["source"]
                    if source == "huggingface":
                        icon = "AI-generated (HuggingFace BART — free, no API key needed)"
                    else:
                        icon = "Extractive fallback"
                    st.caption(icon)
                    st.markdown(f'<div class="summary-card">{summary_data["text"]}</div>', unsafe_allow_html=True)
                    q = summary_data["quality"]
                    st.markdown("**Summary Quality Scores**")
                    st.markdown(
                        render_quality_bar("Coherence",        q["coherence"],   q["coherence_label"]) +
                        render_quality_bar("Readability",       q["readability"], q["readability_label"]) +
                        render_quality_bar("Length Optimality", q["length_opt"],  q["length_label"]) +
                        render_quality_bar("Overall Score",     q["overall"],     q["overall_label"]),
                        unsafe_allow_html=True
                    )
                    cfg = get_section_config(sname)
                    st.caption(f"Target: ~{cfg['target_words']} words | Actual: {len(summary_data['text'].split())} words")
        if other:
            st.markdown("#### Primary Sections")
            for name, content in other.items():
                render_section(name, content)
        if numbered:
            st.markdown("#### Numbered Sections")
            for name, content in sorted(numbered.items(), key=lambda x: float(re.match(r'^(\d+)\.', x[0]).group(1))):
                render_section(name, content)
        if subsect:
            st.markdown("#### Subsections")
            for name, content in sorted(subsect.items(), key=lambda x: [float(n) for n in re.match(r'^(\d+)\.(\d+)', x[0]).groups()]):
                render_section(name, content)
    with tab_viz:
        if writing:
            scores   = writing["scores"]
            stats    = writing["stats"]
            col_r, col_b = st.columns(2)
            with col_r:
                st.markdown("#### Metric Radar")
                cats = ['Language', 'Coherence', 'Reasoning', 'Sophistication', 'Readability']
                vals = [scores[c] for c in cats]
                fig  = go.Figure()
                fig.add_trace(go.Scatterpolar(r=vals, theta=cats, fill='toself',
                                               line_color='#c7522a', fillcolor='rgba(199,82,42,0.2)'))
                fig.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 100])),
                                  showlegend=False, height=360,
                                  paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
                st.plotly_chart(fig, use_container_width=True)
            with col_b:
                st.markdown("#### Text Statistics")
                fig2 = go.Figure(data=[
                    go.Bar(name='Avg Sentence Length', x=['Sentence Length'], y=[stats['avg_sentence_len']], marker_color='#2a5298'),
                    go.Bar(name='Avg Word Length',     x=['Word Length'],     y=[stats['avg_word_len']],     marker_color='#c7522a'),
                ])
                fig2.update_layout(height=360, barmode='group',
                                   paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
                st.plotly_chart(fig2, use_container_width=True)
            st.markdown("#### All Quality Scores")
            fig3 = go.Figure(go.Bar(x=list(scores.keys()), y=list(scores.values()), marker_color='#2a5298'))
            fig3.update_layout(yaxis=dict(range=[0, 100]), height=300,
                               paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
            st.plotly_chart(fig3, use_container_width=True)
        else:
            st.info("Writing quality data not available.")

    with tab_issues:
        st.subheader(" Long Sentences (>30 words)")
        if writing and writing.get("long_sentences"):
            for i, s in enumerate(writing["long_sentences"]):
                st.warning(f"**{i+1}:** {s}")
        else:
            st.success(" No overly long sentences found!")

    with tab_suggestions:
        st.subheader(" Vocabulary Improvements")
        suggestions_map = {
            "very": "extremely", "bad": "adverse",  "good": "beneficial",
            "show": "demonstrate", "big": "substantial", "use": "utilize",
            "help": "facilitate",  "need": "require",    "get": "obtain",
        }
        if writing:
            text_lower = writing["blob"].raw.lower()
            found = False
            for simple, better in suggestions_map.items():
                if simple in text_lower:
                    st.info(f"Consider replacing **'{simple}'** with **'{better}'**")
                    found = True
            if not found:
                st.success(" Great vocabulary! No simple words detected.")
        else:
            st.info("Upload a document to get vocabulary suggestions.")

    with tab_stats:
        st.subheader(" Detailed Text Metrics")
        if writing:
            stats    = writing["stats"]
            dc1, dc2 = st.columns(2)
            with dc1:
                st.write(f"**Words:** {stats['word_count']:,}")
                st.write(f"**Sentences:** {stats['sentence_count']:,}")
                st.write(f"**Avg Sentence Length:** {stats['avg_sentence_len']} words")
            with dc2:
                st.write(f"**Avg Word Length:** {stats['avg_word_len']} chars")
                st.write(f"**Vocabulary Diversity:** {stats['vocab_diversity']}")
                st.write(f"**Complex Word Ratio:** {stats['complex_word_ratio']}")
        else:
            st.info("No stats available.")

    with tab_sentiment:
        st.subheader(" Document Sentiment")
        if writing:
            sentiment = writing["sentiment"]
            st.metric("Sentiment Polarity", sentiment)
            if sentiment > 0.1:
                st.success(" Positive Tone — The paper has an optimistic, affirmative voice.")
            elif sentiment < -0.1:
                st.warning(" Negative/Critical Tone — The paper may use cautionary language.")
            else:
                st.info(" Neutral Tone — The paper maintains an objective, academic tone.")
        else:
            st.info("No sentiment data available.")

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown(f"###  Insights for {st.session_state.user_role}s")
    if st.session_state.user_role == "Student":
        st.success(f"""
         **Learning Insights:**
        - This paper has {len(results['sections'])} sections — read the summaries first for a quick overview
        - Study the methodology to understand research approaches in {results['domain']}
        - The {len(results['keywords'])} extracted keywords are essential concepts to master
        """)
    else:
        st.success(f"""
         **Research Insights:**
        - Comprehensive structure with {len(results['sections'])} sections identified
        - This work contributes to {results['domain']} research
        - Review quality scores to benchmark summary clarity against your own writing
        """)

    st.markdown("<br>", unsafe_allow_html=True)
    col1, col2, col3 = st.columns(3)
    with col1:
        if st.button(" Analyze Another Paper", use_container_width=True):
            for k in ['analysis_complete', 'analysis_results', 'section_summaries', 'paper_summary', 'writing_quality']:
                st.session_state[k] = (False if k == 'analysis_complete' else ({} if k == 'section_summaries' else None))
            st.rerun()
    with col2:
        pdf_bytes = create_pdf_report(
            results['filename'], results, writing, section_summaries, paper_summary or ""
        )
        st.download_button(
            label=" Download PDF Report",
            data=pdf_bytes,
            file_name=f"PaperIQ_Report_{results['filename']}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
    with col3:
        json_results             = {k: v for k, v in results.items() if k != 'sections'}
        json_results['sections'] = list(results['sections'].keys())
        st.download_button(
            label=" Download JSON",
            data=json.dumps(json_results, indent=2),
            file_name=f"paperiq_{results['filename']}.json",
            mime="application/json",
            use_container_width=True,
        )
      
def main():
    initialize_session_state()
    if not st.session_state.authenticated:
        login_page()
    else:
        upload_page()

if __name__ == "__main__":
    main()
